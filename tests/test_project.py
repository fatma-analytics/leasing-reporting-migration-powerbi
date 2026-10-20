from __future__ import annotations

import json
import sqlite3
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "data"
DB = DATA / "leasing_demo.sqlite"


def connection():
    return sqlite3.connect(DB)


def test_manifest_identifies_synthetic_data_and_24_month_history():
    manifest = json.loads((DATA / "manifest.json").read_text(encoding="utf-8"))
    assert manifest["synthetic"] is True
    assert manifest["history_start_month"] == "2024-04"
    assert manifest["history_end_month"] == "2026-03"
    assert manifest["contracts_generated"] == 2000


def test_quality_controls_detect_injected_defects():
    expected = pd.read_csv(DATA / "expected_anomalies.csv").set_index("controle").anomalies_injectees
    observed = pd.read_csv(DATA / "controles_qualite.csv").set_index("controle").anomalies
    assert observed.loc["CONTRATS_DOUBLONS"] == 2 * expected.loc["contrats_doublons"]
    assert observed.loc["CONTRATS_DATES_INCOHERENTES"] == expected.loc["contrats_dates_incoherentes"]
    assert observed.loc["CONTRATS_MONTANTS_NEGATIFS"] == expected.loc["contrats_montants_negatifs"]
    assert observed.loc["ECHEANCES_ORPHELINES"] == expected.loc["echeances_orphelines"]
    assert observed.loc["ECHEANCES_MONTANTS_NEGATIFS"] == expected.loc["echeances_montants_negatifs"]
    raw_payments = pd.read_csv(DATA / "paiements.csv")
    raw_dues = pd.read_csv(DATA / "echeances.csv")[["echeance_id", "montant_du"]]
    joined = raw_payments.merge(raw_dues, on="echeance_id", how="left")
    measured_overpayments = int((joined.montant_paye > joined.montant_du).sum())
    # Some randomly selected negative installments also make their associated payment exceed the due amount.
    assert observed.loc["PAIEMENTS_SUPERIEURS_AU_DU"] == measured_overpayments
    assert measured_overpayments >= expected.loc["paiements_superieurs_au_du"]
    assert observed.loc["PAIEMENTS_ECHEANCE_ABSENTE"] == 0


def test_clean_mart_has_no_key_or_foreign_key_violations():
    con = connection()
    assert con.execute("PRAGMA foreign_key_check").fetchall() == []
    assert con.execute("SELECT COUNT(*) FROM (SELECT contrat_id FROM contrats GROUP BY contrat_id HAVING COUNT(*) > 1)").fetchone()[0] == 0
    assert con.execute("SELECT COUNT(*) FROM contrats WHERE montant_finance <= 0 OR date_fin_prevue < date_debut").fetchone()[0] == 0
    assert con.execute("SELECT COUNT(*) FROM echeances WHERE montant_du <= 0").fetchone()[0] == 0
    assert con.execute("SELECT COUNT(*) FROM paiements p JOIN echeances e ON e.echeance_id=p.echeance_id WHERE p.montant_paye > e.montant_du").fetchone()[0] == 0
    con.close()


def test_arrears_view_has_positive_balances_and_correct_buckets():
    con = connection()
    rows = con.execute("SELECT retard_jours, tranche_retard, reste_a_payer FROM v_impayes").fetchall()
    assert rows
    for days, bucket, remaining in rows:
        expected = "0-30 jours" if days <= 30 else "31-60 jours" if days <= 60 else "61-90 jours" if days <= 90 else ">90 jours"
        assert bucket == expected
        assert days > 0
        assert remaining > 0
    con.close()


def test_aging_case_boundaries():
    con = connection()
    rows = con.execute("""
        WITH boundaries(days) AS (
            SELECT 0 UNION ALL SELECT 30 UNION ALL SELECT 31 UNION ALL
            SELECT 60 UNION ALL SELECT 61 UNION ALL SELECT 90 UNION ALL SELECT 91
        )
        SELECT days,
            CASE WHEN days BETWEEN 0 AND 30 THEN '0-30 jours'
                 WHEN days BETWEEN 31 AND 60 THEN '31-60 jours'
                 WHEN days BETWEEN 61 AND 90 THEN '61-90 jours'
                 ELSE '>90 jours' END
        FROM boundaries ORDER BY days
    """).fetchall()
    assert [bucket for _, bucket in rows] == ["0-30 jours", "0-30 jours", "31-60 jours", "31-60 jours", "61-90 jours", "61-90 jours", ">90 jours"]
    con.close()


def test_acceptance_workbook_contains_15_cases_and_marks_manual_work():
    wb = load_workbook(ROOT / "recette" / "recette_resultats.xlsx", data_only=True)
    ws = wb["Résultats recette"]
    assert ws.max_row == 19  # title rows + header + 15 cases
    statuses = [ws.cell(row, 7).value for row in range(5, 20)]
    assert statuses.count("À valider Desktop") == 2
    assert sum(str(status).startswith("OK") for status in statuses) == 13


def test_legacy_workbooks_and_five_page_previews_exist():
    for filename in [
        "Etat_BO_Production_Mensuelle.xlsx",
        "Reporting_Encours_Excel.xlsx",
        "Suivi_Impayes_Excel.xlsx",
    ]:
        path = ROOT / "legacy_reports" / filename
        assert path.exists() and path.stat().st_size > 10000
        assert load_workbook(path, read_only=True).sheetnames
    encours = load_workbook(ROOT / "legacy_reports" / "Reporting_Encours_Excel.xlsx", data_only=False)
    assert encours["Détail contrats"]["H2"].data_type == "f"
    assert "SUMIFS" in encours["Détail contrats"]["H2"].value
    assert "Paiements contrat" in encours.sheetnames
    previews = sorted((ROOT / "docs" / "screenshots").glob("0[1-5]_*.png"))
    assert len(previews) == 5
    assert all(path.stat().st_size > 20000 for path in previews)
    assert (ROOT / "docs" / "screenshots" / "avant_apres_production.png").exists()
